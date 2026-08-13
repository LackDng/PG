from django.shortcuts import render, redirect
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Sum
from django.contrib import messages
from datetime import date, time, timedelta
from core.models import Invoice, RoomSession, ServiceOrder, OrderItem, ActivityLog
from core.decorators import accountant_required, admin_required, login_required_custom
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def _parse_time(time_str):
    """Parse 'HH:MM' thành time object, trả về None nếu không hợp lệ."""
    try:
        h, m = time_str.strip().split(":")
        return time(int(h), int(m))
    except Exception:
        return None


def _get_report_data(dt_from, dt_to):
    """Lấy dữ liệu báo cáo trong khoảng dt_from..dt_to (aware datetimes)."""

    invoices = Invoice.objects.filter(
        created_at__gte=dt_from,
        created_at__lt=dt_to,
    ).select_related("session__room", "created_by").order_by("-created_at")

    agg = invoices.aggregate(
        s_service=Sum("total_service"),
        s_food=Sum("total_food"),
        s_outside=Sum("total_outside"),
        s_discount=Sum("discount_amount"),
        s_revenue=Sum("total_after_discount"),
        s_tip=Sum("tip"),
        s_cash=Sum("cash_amount"),
        s_transfer=Sum("transfer_amount"),
    )
    total_service  = agg["s_service"]  or 0
    total_food     = agg["s_food"]     or 0
    total_outside  = agg["s_outside"]  or 0
    total_discount = agg["s_discount"] or 0
    total_revenue  = agg["s_revenue"]  or 0
    total_tip      = agg["s_tip"]      or 0
    total_cash     = agg["s_cash"]     or 0
    total_transfer = agg["s_transfer"] or 0

    room_summaries = {}
    for inv in invoices:
        rname = inv.session.room.name
        if rname not in room_summaries:
            room_summaries[rname] = {"room_name": rname, "count": 0, "service": 0, "food": 0, "outside": 0, "total": 0}
        room_summaries[rname]["count"]   += 1
        room_summaries[rname]["service"] += inv.total_service
        room_summaries[rname]["food"]    += inv.total_food
        room_summaries[rname]["outside"] += inv.total_outside
        room_summaries[rname]["total"]   += inv.total_after_discount + inv.tip

    room_summaries = sorted(room_summaries.values(), key=lambda x: x["total"], reverse=True)

    return {
        "invoices":       invoices,
        "total_service":  total_service,
        "total_food":     total_food,
        "total_outside":  total_outside,
        "total_discount": total_discount,
        "total_revenue":  total_revenue,
        "total_tip":      total_tip,
        "total_cash":     total_cash,
        "total_transfer": total_transfer,
        "room_summaries": room_summaries,
        "invoice_count":  invoices.count(),
    }


def _parse_date(date_str, fallback):
    try:
        return date.fromisoformat(date_str)
    except Exception:
        return fallback


def _build_range(date_from, time_from_str, date_to, time_to_str):
    """Tạo cặp aware datetime từ các tham số lọc."""
    tz = timezone.get_current_timezone()
    t_from = _parse_time(time_from_str) or time(0, 0)
    dt_from = timezone.datetime.combine(date_from, t_from).replace(tzinfo=tz)

    t_to = _parse_time(time_to_str)
    if t_to:
        dt_to = timezone.datetime.combine(date_to, t_to).replace(tzinfo=tz)
    else:
        dt_to = timezone.datetime.combine(date_to, time(0, 0)).replace(tzinfo=tz) + timedelta(days=1)
    return dt_from, dt_to


@accountant_required
def daily_report(request):
    today = timezone.localdate()
    date_from_str = request.GET.get("date_from", today.strftime("%Y-%m-%d"))
    date_to_str   = request.GET.get("date_to",   today.strftime("%Y-%m-%d"))
    time_from_str = request.GET.get("time_from", "")
    time_to_str   = request.GET.get("time_to",   "")

    date_from = _parse_date(date_from_str, today)
    date_to   = _parse_date(date_to_str,   today)
    if date_from > date_to:
        date_to = date_from

    dt_from, dt_to = _build_range(date_from, time_from_str, date_to, time_to_str)
    data = _get_report_data(dt_from, dt_to)
    return render(request, "reports/daily.html", {
        "date_from_str": date_from.strftime("%Y-%m-%d"),
        "date_to_str":   date_to.strftime("%Y-%m-%d"),
        "time_from_str": time_from_str,
        "time_to_str":   time_to_str,
        "dt_from":       dt_from,
        "dt_to":         dt_to,
        **data,
    })


@accountant_required
def export_excel(request):
    """Xuất báo cáo ra file Excel theo khoảng ngày+giờ."""
    today = timezone.localdate()
    date_from_str = request.GET.get("date_from", today.strftime("%Y-%m-%d"))
    date_to_str   = request.GET.get("date_to",   today.strftime("%Y-%m-%d"))
    time_from_str = request.GET.get("time_from", "")
    time_to_str   = request.GET.get("time_to",   "")

    date_from = _parse_date(date_from_str, today)
    date_to   = _parse_date(date_to_str,   today)
    if date_from > date_to:
        date_to = date_from

    dt_from, dt_to = _build_range(date_from, time_from_str, date_to, time_to_str)
    data = _get_report_data(dt_from, dt_to)
    wb = _build_excel(dt_from, dt_to, data)

    tf = time_from_str.replace(":", "") if time_from_str else ""
    tt = time_to_str.replace(":", "")   if time_to_str   else ""
    time_suffix = f"_{tf}-{tt}" if (tf or tt) else ""
    if date_from == date_to:
        filename = f"baocao_{date_from.strftime('%Y%m%d')}{time_suffix}.xlsx"
    else:
        filename = f"baocao_{date_from.strftime('%Y%m%d')}_den_{date_to.strftime('%Y%m%d')}{time_suffix}.xlsx"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def _build_excel(dt_from, dt_to, data):
    """Tạo workbook Excel từ dữ liệu báo cáo."""
    wb = openpyxl.Workbook()

    # ─── Sheet 1: Tổng quan ───────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Tổng quan"

    header_font    = Font(bold=True, size=12)
    header_fill    = PatternFill("solid", fgColor="1F3864")
    subheader_fill = PatternFill("solid", fgColor="2E75B6")
    white_font     = Font(bold=True, color="FFFFFF", size=11)
    center         = Alignment(horizontal="center", vertical="center")
    right          = Alignment(horizontal="right", vertical="center")
    thin           = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # Tiêu đề
    local_from = timezone.localtime(dt_from)
    local_to   = timezone.localtime(dt_to)
    date_from_d = local_from.date()
    date_to_d   = (local_to - timedelta(seconds=1)).date()
    if date_from_d == date_to_d:
        date_label = date_from_d.strftime("%d/%m/%Y")
    else:
        date_label = f"{date_from_d.strftime('%d/%m/%Y')} – {date_to_d.strftime('%d/%m/%Y')}"
    time_label = f"  {local_from.strftime('%H:%M')} – {local_to.strftime('%H:%M')}"
    ws1.merge_cells("A1:F1")
    ws1["A1"] = f"BÁO CÁO DOANH THU  {date_label}  {time_label}"
    ws1["A1"].font  = Font(bold=True, size=16, color="1F3864")
    ws1["A1"].alignment = center

    ws1.merge_cells("A2:F2")
    ws1["A2"] = f"Tổng số hóa đơn: {data['invoice_count']}"
    ws1["A2"].alignment = center

    # Bảng tổng hợp
    ws1["A4"] = "CHỈ TIÊU"
    ws1["B4"] = "GIÁ TRỊ (VNĐ)"
    for cell in [ws1["A4"], ws1["B4"]]:
        cell.font = white_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    summary_rows = [
        ("Tổng doanh thu",       data["total_revenue"]),
        ("Dịch vụ",              data["total_service"]),
        ("Đồ ăn/uống",           data["total_food"]),
        ("Mua ngoài",            data["total_outside"]),
        ("Tổng giảm giá",        -data["total_discount"]),
        ("Tip",                  data["total_tip"]),
        ("",                     None),
        ("Tiền mặt",             data["total_cash"]),
        ("Chuyển khoản",         data["total_transfer"]),
    ]
    for i, (label, val) in enumerate(summary_rows, start=5):
        ws1.cell(row=i, column=1, value=label).border = thin
        cell = ws1.cell(row=i, column=2, value=val)
        cell.border = thin
        cell.alignment = right
        if val is not None:
            cell.number_format = '#,##0'
        if label == "Tổng doanh thu":
            ws1.cell(row=i, column=1).font = Font(bold=True)
            cell.font = Font(bold=True)

    # Bảng theo phòng
    row = 5 + len(summary_rows) + 2
    ws1.cell(row=row, column=1, value="DOANH THU THEO PHÒNG").font = header_font
    row += 1
    room_headers = ["Phòng", "Số lần", "Dịch vụ", "Đồ ăn/uống", "Mua ngoài", "Tổng"]
    for c, h in enumerate(room_headers, 1):
        cell = ws1.cell(row=row, column=c, value=h)
        cell.font = white_font
        cell.fill = subheader_fill
        cell.alignment = center
        cell.border = thin
    row += 1
    for rs in data["room_summaries"]:
        vals = [rs["room_name"], rs["count"], rs["service"], rs["food"], rs["outside"], rs["total"]]
        for c, v in enumerate(vals, 1):
            cell = ws1.cell(row=row, column=c, value=v)
            cell.border = thin
            if c >= 3:
                cell.number_format = '#,##0'
                cell.alignment = right
        row += 1

    ws1.column_dimensions["A"].width = 25
    ws1.column_dimensions["B"].width = 18
    for col in "CDEF":
        ws1.column_dimensions[col].width = 18

    # ─── Sheet 2: Chi tiết hóa đơn ────────────────────────────────────────────
    ws2 = wb.create_sheet("Chi tiết hóa đơn")

    detail_headers = [
        "#", "Phòng", "Giờ thanh toán",
        "Dịch vụ (đ)", "Đồ ăn (đ)", "Mua ngoài (đ)",
        "Giảm giá %", "Giảm giá (đ)", "Sau giảm giá",
        "Tip (đ)", "Tổng (đ)", "Hình thức TT",
        "Tiền mặt (đ)", "Chuyển khoản (đ)", "Thu ngân",
    ]
    for c, h in enumerate(detail_headers, 1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.font = white_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    pm_labels = {"cash": "Tiền mặt", "transfer": "Chuyển khoản", "mixed": "Kết hợp"}
    for r, inv in enumerate(data["invoices"], start=2):
        row_data = [
            inv.id,
            inv.session.room.name,
            timezone.localtime(inv.created_at).strftime("%d/%m/%Y %H:%M"),
            inv.total_service,
            inv.total_food,
            inv.total_outside,
            inv.discount_percent,
            inv.discount_amount,
            inv.total_after_discount,
            inv.tip,
            inv.total_after_discount + inv.tip,
            pm_labels.get(inv.payment_method, inv.payment_method),
            inv.cash_amount,
            inv.transfer_amount,
            inv.created_by.get_display_name() if inv.created_by else "",
        ]
        money_cols = {4, 5, 6, 8, 9, 10, 11, 13, 14}
        for c, v in enumerate(row_data, 1):
            cell = ws2.cell(row=r, column=c, value=v)
            cell.border = thin
            if c in money_cols and isinstance(v, (int, float)):
                cell.number_format = '#,##0'
                cell.alignment = right

    col_widths = [6, 15, 20, 15, 15, 15, 12, 15, 15, 12, 15, 16, 16, 18, 18]
    for i, w in enumerate(col_widths, 1):
        ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = w

    # ─── Sheet 3: Chi tiết từng bill ──────────────────────────────────────────
    ws3 = wb.create_sheet("Chi tiết bill")

    detail3_headers = [
        "HĐ #", "Phòng", "Giờ TT",
        "Loại", "Tên hàng", "Bắt đầu", "Kết thúc", "Thời gian",
        "SL", "Đơn giá (đ)", "Thành tiền (đ)",
    ]
    for c, h in enumerate(detail3_headers, 1):
        cell = ws3.cell(row=1, column=c, value=h)
        cell.font = white_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    money_cols3 = {10, 11}
    r3 = 2
    for inv in data["invoices"]:
        session = inv.session
        inv_time = timezone.localtime(inv.created_at).strftime("%d/%m/%Y %H:%M")

        for so in session.get_all_service_orders().exclude(status=ServiceOrder.STATUS_CANCELLED).order_by("started_at"):
            end = so.ended_at or inv.created_at
            dur_min = int((end - so.started_at).total_seconds() / 60)
            dur_str = f"{dur_min // 60}h{dur_min % 60:02d}p" if dur_min >= 60 else f"{dur_min}p"
            cost = so.calculate_cost(at_time=inv.created_at)
            row3 = [
                inv.id, session.room.name, inv_time,
                "Dịch vụ", so.service_name,
                timezone.localtime(so.started_at).strftime("%H:%M"),
                timezone.localtime(end).strftime("%H:%M"),
                dur_str, "", so.base_price, cost,
            ]
            for c, v in enumerate(row3, 1):
                cell = ws3.cell(row=r3, column=c, value=v)
                cell.border = thin
                if c in money_cols3 and isinstance(v, (int, float)):
                    cell.number_format = '#,##0'
                    cell.alignment = right
            r3 += 1

        for item in session.get_all_order_items().order_by("created_at"):
            row3 = [
                inv.id, session.room.name, inv_time,
                "Đồ ăn/uống" if item.item_type == "menu" else "Mua ngoài",
                item.name, "", "", "",
                item.quantity, item.unit_price, item.subtotal,
            ]
            for c, v in enumerate(row3, 1):
                cell = ws3.cell(row=r3, column=c, value=v)
                cell.border = thin
                if c in money_cols3 and isinstance(v, (int, float)):
                    cell.number_format = '#,##0'
                    cell.alignment = right
            r3 += 1

    col3_widths = [8, 15, 18, 14, 25, 10, 10, 12, 6, 16, 16]
    for i, w in enumerate(col3_widths, 1):
        ws3.column_dimensions[ws3.cell(row=1, column=i).column_letter].width = w

    return wb


@admin_required
def delete_revenue(request):
    """Xóa dữ liệu doanh thu theo khoảng thời gian (admin only)."""
    if request.method == "POST":
        date_from_str = request.POST.get("date_from", "")
        date_to_str   = request.POST.get("date_to", "")
        confirm       = request.POST.get("confirm") == "yes"

        if not confirm:
            messages.error(request, "Vui lòng tick xác nhận trước khi xóa.")
            return redirect("delete_revenue")

        try:
            date_from = date.fromisoformat(date_from_str)
            date_to   = date.fromisoformat(date_to_str)
        except ValueError:
            messages.error(request, "Ngày không hợp lệ.")
            return redirect("delete_revenue")

        if date_from > date_to:
            messages.error(request, "Ngày bắt đầu phải trước ngày kết thúc.")
            return redirect("delete_revenue")

        tz = timezone.get_current_timezone()
        dt_from = timezone.datetime.combine(date_from, timezone.datetime.min.time()).replace(tzinfo=tz)
        dt_to   = timezone.datetime.combine(date_to, timezone.datetime.max.time()).replace(tzinfo=tz)

        invoices = Invoice.objects.filter(created_at__gte=dt_from, created_at__lte=dt_to)
        count = invoices.count()
        # Xóa invoice trước (cascade sẽ không xóa session), chỉ xóa invoice
        invoices.delete()

        messages.success(request, f"Đã xóa {count} hóa đơn từ {date_from_str} đến {date_to_str}.")
        return redirect("delete_revenue")

    return render(request, "reports/delete_revenue.html")


@login_required_custom
def activity_log_list(request):
    from core.decorators import login_required_custom as _
    logs = ActivityLog.objects.select_related("user", "session__room").order_by("-created_at")

    today = timezone.localdate().strftime("%Y-%m-%d")
    date_str = request.GET.get("date", today)
    action_filter = request.GET.get("action", "")

    try:
        filter_date = date.fromisoformat(date_str)
        tz = timezone.get_current_timezone()
        day_start = timezone.datetime.combine(filter_date, timezone.datetime.min.time()).replace(tzinfo=tz)
        day_end = day_start + timedelta(days=1)
        logs = logs.filter(created_at__gte=day_start, created_at__lt=day_end)
    except ValueError:
        pass

    if action_filter:
        logs = logs.filter(action=action_filter)

    return render(request, "reports/activity_logs.html", {
        "logs": logs[:500],
        "action_choices": ActivityLog.ACTION_CHOICES,
        "filter_date": date_str,
        "filter_action": action_filter,
    })
